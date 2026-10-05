#!/usr/bin/env python3
"""Builds the Unit 5 report as a Word document from the Markdown sources.

Needs python-docx:
    python3 build_report.py

Formatting follows the assignment brief: Times New Roman 12 pt, 1.5 line
spacing, left-aligned paragraphs, A4 portrait, "Page X of Y" in the footer, a
table of contents, Harvard references and the standard Learn Key front page
carrying the word count. Code listings use Courier New, as is conventional.

The table of contents is written with its page numbers already filled in.
Those numbers come from toc_pages.json, which measure_pages.py produces from a
real render of the document, so the contents page is correct when first opened.
"""
from __future__ import annotations

import copy
import json
import re
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
WORK = ROOT.parent
CONTENT = ROOT / "content"
LOGOS = ROOT / "logos"
OUTPUT = WORK / "out"
DOCX_PATH = OUTPUT / "MaltaShop_Microservices_Report.docx"
REPOSITORY = "https://github.com/imranneta5555/maltashop-microservices"
TOC_PAGES = ROOT / "toc_pages.json"
COVER = json.loads((ROOT / "cover.json").read_text(encoding="utf-8"))

BODY_FONT = "Times New Roman"
CODE_FONT = "Courier New"
CODE_SIZE = 9          # short snippets in the text (83 characters per line)
LISTING_SIZE = 7.5     # whole files (100 characters per line)
KEEP_TOGETHER_LINES = 30
BLACK = RGBColor(0, 0, 0)
MARGIN_CM = 2.54
COVER_MARGIN_CM = 1.8       # the front page keeps the wider layout of the official form
COVER_WIDTH_CM = 21.0 - 2 * COVER_MARGIN_CM
TEXT_WIDTH_CM = 21.0 - 2 * MARGIN_CM
DESIGN_WIDTH = 880          # width, in SVG units, that the diagrams were drawn at

APPENDIX_SOURCES = [
    ("B.1 docker-compose.yml", "docker-compose.yml"),
    ("B.2 order-service/app/main.py", "order-service/app/main.py"),
    ("B.3 order-service/app/db.py", "order-service/app/db.py"),
    ("B.4 order-service/app/outbox.py", "order-service/app/outbox.py"),
    ("B.5 notification-service/app/consumer.py", "notification-service/app/consumer.py"),
]


# ============================================================ Markdown model

@dataclass
class Block:
    kind: str
    text: str = ""
    items: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    target: str = ""


BLOCK_START = re.compile(r"^(#{1,3} |!\[|Table: |Listing: |@file |\||```|> |- |\d+\. )")


def parse_markdown(source: str) -> list[Block]:
    """Parses the small subset of Markdown the report content uses."""
    lines = source.splitlines()
    blocks: list[Block] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code = []
            while not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(Block("code", "\n".join(code)))
            i += 1
            continue
        heading = re.match(r"^(#{1,3}) (.+)$", line)
        if heading:
            blocks.append(Block(f"h{len(heading.group(1))}", heading.group(2).strip()))
            i += 1
            continue
        figure = re.match(r"^!\[(.+)\]\((.+?)\)(?:\{w=([\d.]+)\})?\s*$", line)
        if figure:
            blocks.append(Block("figure", figure.group(1), target=figure.group(2),
                                items=[float(figure.group(3))] if figure.group(3) else []))
            i += 1
            continue
        if line.startswith("Table: "):
            blocks.append(Block("table_caption", line[len("Table: "):]))
            i += 1
            continue
        if line.startswith("Listing: "):
            blocks.append(Block("listing_caption", line[len("Listing: "):]))
            i += 1
            continue
        if line.startswith("@file "):
            path = line[len("@file "):].strip()
            blocks.append(Block("listing", (WORK / path).read_text(encoding="utf-8").rstrip("\n"), target=path))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(Block("table", rows=rows))
            continue
        if line.startswith("> "):
            blocks.append(Block("quote", line[2:]))
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            items = []
            while i < len(lines) and re.match(r"^\d+\. ", lines[i]):
                items.append(re.sub(r"^\d+\. ", "", lines[i]))
                i += 1
            blocks.append(Block("numbered", items=items))
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:])
                i += 1
            blocks.append(Block("bullets", items=items))
            continue
        paragraph = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not BLOCK_START.match(lines[i]):
            paragraph.append(lines[i])
            i += 1
        blocks.append(Block("paragraph", " ".join(paragraph)))
    return blocks


INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


def inline_runs(text: str) -> list[tuple[str, str]]:
    """Splits text into (style, content) pieces: plain, code, bold or italic."""
    pieces = []
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`"):
            pieces.append(("code", part[1:-1]))
        elif part.startswith("**"):
            pieces.append(("bold", part[2:-2]))
        elif part.startswith("*") and len(part) > 1:
            pieces.append(("italic", part[1:-1]))
        else:
            pieces.append(("plain", part))
    return pieces


# ============================================================ low-level XML

def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), name)


def configure_style(style, name: str, size: float, bold=None, italic=None) -> None:
    """Sets a style's font fully, removing the theme fonts and colours Word would otherwise use."""
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if fonts.get(qn(attribute)) is not None:
            del fonts.attrib[qn(attribute)]
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), name)
    colour = rpr.find(qn("w:color"))
    if colour is not None:
        rpr.remove(colour)
    style.font.color.rgb = BLACK
    language = rpr.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        rpr.append(language)
    language.set(qn("w:val"), "en-GB")


def _shading(fill: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    return shading


# Elements that must follow w:shd inside w:pPr and w:tcPr (ECMA-376 sequence order).
PPR_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                 "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                 "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                 "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                 "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
                 "w:sectPr", "w:pPrChange")
TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                  "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")


def shade_paragraph(paragraph, fill: str) -> None:
    paragraph._p.get_or_add_pPr().insert_element_before(_shading(fill), *PPR_AFTER_SHD)


def shade_cell(cell, fill: str) -> None:
    cell._tc.get_or_add_tcPr().insert_element_before(_shading(fill), *TCPR_AFTER_SHD)


def add_field(paragraph, instruction: str, cached: str) -> None:
    """Adds a complex field (e.g. PAGE) with a cached result."""
    for kind, payload in (("begin", None), ("instr", instruction), ("separate", None),
                          ("result", cached), ("end", None)):
        run = paragraph.add_run()
        set_run_font(run, BODY_FONT, 12)
        if kind == "instr":
            text = OxmlElement("w:instrText")
            text.set(qn("xml:space"), "preserve")
            text.text = f" {payload} "
            run._element.append(text)
        elif kind == "result":
            run.text = payload
        else:
            char = OxmlElement("w:fldChar")
            char.set(qn("w:fldCharType"), kind)
            run._element.append(char)


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


def cell_borders(cell, edges=("top", "bottom", "left", "right"), size=6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        if edge in edges:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), str(size))
            element.set(qn("w:color"), "000000")
        else:
            element.set(qn("w:val"), "nil")
        borders.append(element)
    tc_pr.append(borders)


def set_column_widths(table, widths_cm: list[float]) -> None:
    """Writes the widths into the table grid as well as every cell.

    Word reads the cell widths, but LibreOffice and Google Docs lay the table
    out from the grid, which python-docx otherwise leaves split evenly."""
    grid = table._tbl.tblGrid
    for column, width in zip(grid.findall(qn("w:gridCol")), widths_cm):
        column.set(qn("w:w"), str(Cm(width).twips))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
    table_width = table._tbl.tblPr.find(qn("w:tblW"))
    if table_width is not None:
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(Cm(sum(widths_cm)).twips))


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_as_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    tr_pr.append(header)


# ============================================================ document setup

def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(MARGIN_CM))
    section.footer_distance = Cm(1.25)

    styles = doc.styles
    normal = styles["Normal"]
    configure_style(normal, BODY_FONT, 12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for level, size, before, after in ((1, 16, 0, 12), (2, 14, 12, 6), (3, 12, 10, 4)):
        style = styles[f"Heading {level}"]
        configure_style(style, BODY_FONT, size, bold=True, italic=False)
        pf = style.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.5
        pf.keep_with_next = True
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.page_break_before = level == 1

    caption = styles["Caption"]
    configure_style(caption, BODY_FONT, 12, bold=False, italic=True)
    caption.paragraph_format.line_spacing = 1.5
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(12)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for name in ("List Bullet", "List Number", "Quote"):
        style = styles[name]
        configure_style(style, BODY_FONT, 12, italic=name == "Quote")
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    styles["Quote"].paragraph_format.left_indent = Cm(1.0)
    styles["Quote"].paragraph_format.right_indent = Cm(1.0)

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = normal
    configure_style(code, CODE_FONT, CODE_SIZE)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.left_indent = Cm(0.6)          # wrapped lines hang under the first
    code.paragraph_format.first_line_indent = Cm(-0.4)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    listing = styles.add_style("Code Listing", WD_STYLE_TYPE.PARAGRAPH)
    listing.base_style = code
    configure_style(listing, CODE_FONT, LISTING_SIZE)

    table_text = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    table_text.base_style = normal
    configure_style(table_text, BODY_FONT, 12)
    table_text.paragraph_format.line_spacing = 1.0
    table_text.paragraph_format.space_before = Pt(2)
    table_text.paragraph_format.space_after = Pt(2)

    contents_title = styles.add_style("Contents Title", WD_STYLE_TYPE.PARAGRAPH)
    contents_title.base_style = normal
    configure_style(contents_title, BODY_FONT, 16, bold=True)
    contents_title.paragraph_format.space_after = Pt(12)

    for level in (1, 2):
        style = styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        configure_style(style, BODY_FONT, 12, bold=level == 1)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.left_indent = Cm(0 if level == 1 else 0.8)
        style.paragraph_format.tab_stops.add_tab_stop(
            Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    core = doc.core_properties
    core.title = COVER["assignment_title"]
    core.subject = COVER["unit_title"]
    core.author = COVER["learner_name"]
    core.last_modified_by = COVER["learner_name"]
    core.comments = ""                      # python-docx's template says "generated by python-docx"
    core.revision = 1
    core.created = core.modified = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    core.language = "en-GB"

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for piece in ("Page ", None, " of ", "NUMPAGES"):
        if piece is None:
            add_field(footer, "PAGE", "1")
        elif piece == "NUMPAGES":
            add_field(footer, "NUMPAGES", "1")
        else:
            run = footer.add_run(piece)
            set_run_font(run, BODY_FONT, 12)
    return doc


# ============================================================ writers

def write_inline(paragraph, text: str, size: float = 12, code_size: float = 11) -> None:
    for style, content in inline_runs(text):
        if style == "code":
            run = paragraph.add_run()
            # A no-break hyphen stops "P->Link" being split across two lines.
            for index, piece in enumerate(content.split("->")):
                if index:
                    run._r.append(OxmlElement("w:noBreakHyphen"))
                    run._r.add_t(">")
                if piece:
                    run._r.add_t(piece)
            set_run_font(run, CODE_FONT, code_size)
        else:
            run = paragraph.add_run(content)
            set_run_font(run, BODY_FONT, size)
            run.bold = style == "bold"
            run.italic = style == "italic"


def caption_paragraph(doc, text: str, keep_with_next: bool) -> None:
    label = re.match(r"^((?:Figure|Table|Listing) \d+)\s+—\s+(.*)$", text)
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.paragraph_format.keep_with_next = keep_with_next
    if label:
        bold = paragraph.add_run(f"{label.group(1)}: ")
        set_run_font(bold, BODY_FONT, 12)
        bold.bold = True
        write_inline(paragraph, label.group(2))
        for run in paragraph.runs[1:]:
            run.italic = True
    else:
        write_inline(paragraph, text)


def write_code(doc, source: str, style: str = "Code Block") -> None:
    lines = source.replace("\t", "    ").splitlines() or [""]
    keep = len(lines) <= KEEP_TOGETHER_LINES
    for index, line in enumerate(lines):
        paragraph = doc.add_paragraph(style=style)
        run = paragraph.add_run(line if line else " ")
        set_run_font(run, CODE_FONT, CODE_SIZE if style == "Code Block" else LISTING_SIZE)
        shade_paragraph(paragraph, "F2F2F2")
        # Short blocks stay on one page; long listings may break across pages.
        paragraph.paragraph_format.keep_with_next = keep and index < len(lines) - 1


def write_figure(doc, caption: str, target: str, width: float | None = None) -> None:
    image = WORK / target
    if not image.exists():
        raise SystemExit(f"missing figure image: {image}")
    if width is None:  # a diagram: scale from its design width
        pixel_width = int(subprocess.run(
            ["sips", "-g", "pixelWidth", str(image)], capture_output=True, text=True, check=True
        ).stdout.split()[-1])
        width = (pixel_width / 3) / DESIGN_WIDTH * TEXT_WIDTH_CM
    width_cm = min(TEXT_WIDTH_CM, width)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.add_run().add_picture(str(image), width=Cm(width_cm))
    caption_paragraph(doc, caption, keep_with_next=False)


CODE_CHAR_CM = 0.212      # Courier New 10 pt is 6 pt per character
CELL_PADDING_CM = 0.45
PT_TO_CM = 2.54 / 72


def text_width_cm(word: str, bold: bool) -> float:
    """Generous estimate of a word's width in Times New Roman 12 pt."""
    em = 0.0
    for char in word:
        if char.isupper():
            em += 0.72
        elif char.isdigit() or char == "_":
            em += 0.50
        elif char.isalpha():
            em += 0.48
        else:
            em += 0.34
    return em * 12 * PT_TO_CM * (1.08 if bold else 1.0)


def unbreakable_width(text: str, header: bool) -> float:
    """Width needed by the longest piece of a cell that should not be split."""
    widest = 0.0
    for style, content in inline_runs(text):
        if style == "code":
            pieces = [content] if len(content) <= 18 else content.split()
            widths = [len(piece) * CODE_CHAR_CM for piece in pieces]
        else:
            bold = header or style == "bold"
            widths = [text_width_cm(piece, bold) for piece in content.split()]
        widest = max([widest] + widths)
    return widest + CELL_PADDING_CM


def column_widths(rows: list[list[str]]) -> list[float]:
    """Shares the text width by content length, but never so narrowly that a word breaks."""
    columns = len(rows[0])
    weights, minimum = [], []
    for column in range(columns):
        cells = [(index, row[column]) for index, row in enumerate(rows) if column < len(row)]
        longest = max(len(re.sub(r"[`*]", "", text)) for _, text in cells)
        weights.append(min(max(longest, 6), 58))
        minimum.append(max(unbreakable_width(text, header=index == 0) for index, text in cells))
    if sum(minimum) >= TEXT_WIDTH_CM:
        return [TEXT_WIDTH_CM * width / sum(minimum) for width in minimum]

    widths = [TEXT_WIDTH_CM * weight / sum(weights) for weight in weights]
    for _ in range(columns):
        short = [c for c in range(columns) if widths[c] < minimum[c] - 1e-9]
        if not short:
            break
        deficit = sum(minimum[c] - widths[c] for c in short)
        for c in short:
            widths[c] = minimum[c]
        spare = {c: widths[c] - minimum[c] for c in range(columns) if widths[c] > minimum[c]}
        room = sum(spare.values())
        for c, amount in spare.items():
            widths[c] -= deficit * amount / room
    return widths


def write_table(doc, rows: list[list[str]]) -> None:
    widths = column_widths(rows)
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.style = doc.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False  # writes a fixed tblLayout in the schema-correct position

    for row_index, cells in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        if row_index == 0:
            repeat_as_header(row)
        for column, content in enumerate(cells):
            cell = row.cells[column]
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles["Table Text"]
            write_inline(paragraph, content, size=12, code_size=10)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
                shade_cell(cell, "E7E6E6")
    set_column_widths(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    spacer.paragraph_format.line_spacing = 1.0


def new_numbering(doc) -> int:
    """A new numbering instance for the List Number style, restarting at 1."""
    numbering = doc.part.numbering_part.element
    style_num = doc.styles["List Number"].element.pPr.numPr.numId.val
    abstract_id = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract_id)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return num.numId


def write_blocks(doc, blocks: list[Block], bookmarks: dict[str, str], counter: list[int]) -> None:
    pending_table_caption = None
    for block in blocks:
        if block.kind in ("h1", "h2", "h3"):
            level = int(block.kind[1])
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            write_inline(paragraph, block.text, size={1: 16, 2: 14, 3: 12}[level])
            for run in paragraph.runs:
                run.bold = True
            if block.text in bookmarks:
                counter[0] += 1
                add_bookmark(paragraph, bookmarks[block.text], counter[0])
        elif block.kind == "paragraph":
            write_inline(doc.add_paragraph(), block.text)
        elif block.kind == "quote":
            write_inline(doc.add_paragraph(style="Quote"), block.text)
        elif block.kind == "bullets":
            for item in block.items:
                write_inline(doc.add_paragraph(style="List Bullet"), item)
        elif block.kind == "numbered":
            num_id = new_numbering(doc)  # every numbered list starts again at 1
            for item in block.items:
                paragraph = doc.add_paragraph(style="List Number")
                num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                num_pr.get_or_add_ilvl().val = 0
                num_pr.get_or_add_numId().val = num_id
                write_inline(paragraph, item)
        elif block.kind == "code":
            write_code(doc, block.text)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif block.kind == "figure":
            write_figure(doc, block.text, block.target, block.items[0] if block.items else None)
        elif block.kind == "listing_caption":
            caption_paragraph(doc, block.text, keep_with_next=True)
        elif block.kind == "listing":
            write_code(doc, block.text, style="Code Listing")
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif block.kind == "table_caption":
            pending_table_caption = block.text
        elif block.kind == "table":
            if pending_table_caption:
                caption_paragraph(doc, pending_table_caption, keep_with_next=True)
                pending_table_caption = None
            write_table(doc, block.rows)


# ============================================================ front matter

def cover_sheet(doc, word_count: int) -> None:
    """The standard Learn Key front page, laid out as on the previous (DevOps) submission."""
    size = 10.5

    def text_run(paragraph, text, points=size, bold=False, italic=False):
        run = paragraph.add_run(text)
        set_run_font(run, BODY_FONT, points)
        run.bold, run.italic = bold, italic
        return run

    def tight(paragraph, before=0, after=0):
        paragraph.style = doc.styles["Table Text"]
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        return paragraph

    # Header: logo, title, and the ruled OTHM / MFHEA box on the right.
    head = doc.add_table(rows=1, cols=3)
    head.autofit = False
    cells = head.rows[0].cells
    for cell in cells:
        cell_borders(cell, edges=())
    set_column_widths(head, [4.5, 6.3, 6.6])

    tight(cells[0].paragraphs[0]).add_run().add_picture(str(LOGOS / "learnkey.png"), height=Cm(1.42))
    text_run(tight(cells[1].paragraphs[0], before=8), "Assignment Cover Sheet", points=14, bold=True)

    accreditation = cells[2].add_table(rows=0, cols=2)
    accreditation.autofit = False
    for filename, height in (("othm.png", 1.2), ("mfhea.png", 1.0)):
        row = accreditation.add_row()
        left, right = row.cells
        tight(left.paragraphs[0], before=4, after=4).add_run().add_picture(str(LOGOS / filename), height=Cm(height))
        tight(right.paragraphs[0])
        cell_borders(left)
        cell_borders(right)
    set_column_widths(accreditation, [4.95, 1.27])
    tight(cells[2].paragraphs[0])
    tight(cells[2].paragraphs[-1])  # Word requires a paragraph after a nested table

    instruction = doc.add_paragraph()
    instruction.paragraph_format.line_spacing = 1.0
    instruction.paragraph_format.space_before = Pt(16)
    instruction.paragraph_format.space_after = Pt(8)
    text_run(instruction, "This cover sheet must be completed and added to the front of every assignment", bold=True)

    fields = [
        ("Learner Name and Surname:", COVER["learner_name"], False),
        ("Learner Registration No.", COVER["registration_no"], False),
        ("Study Centre Name", COVER["study_centre"], True),
        ("Qualification Title", COVER["qualification"], False),
        ("Unit Reference No.", COVER["unit_reference"], False),
        ("Unit Title", COVER["unit_title"], False),
        ("Assignment Title", COVER["assignment_title"], False),
        ("Submission Date", COVER["submission_date"], False),
        ("Word Count", f"{word_count:,} words (excluding the front page, contents, figures, captions, "
                       "tables, code, references and appendices)", False),
    ]
    form = doc.add_table(rows=0, cols=2)
    form.autofit = False
    for label, value, bold_value in fields:
        row = form.add_row()
        text_run(tight(row.cells[0].paragraphs[0], before=3.5, after=3.5), label, bold=True, italic=True)
        text_run(tight(row.cells[1].paragraphs[0], before=3.5, after=3.5), value, bold=bold_value)
        cell_borders(row.cells[0])
        cell_borders(row.cells[1])
    set_column_widths(form, [5.75, COVER_WIDTH_CM - 5.75])

    declaration_table = doc.add_table(rows=1, cols=1)
    set_column_widths(declaration_table, [COVER_WIDTH_CM])
    declaration = declaration_table.rows[0].cells[0]
    text_run(tight(declaration.paragraphs[0], before=7, after=6), "Declaration of authenticity:", bold=True)
    statements = [
        "I declare that the attached submission is my own original work. No significant part of it has "
        "been submitted for any other assignment and I have acknowledged in my notes and bibliography all "
        "written and electronic sources used.",
        "I acknowledge that my assignment will be subject to electronic scrutiny for academic honesty.",
        "I understand that failure to meet these guidelines may instigate the centre's malpractice "
        "procedures and risk failure of the unit and / or qualification.",
    ]
    for number, statement in enumerate(statements, start=1):
        paragraph = tight(declaration.add_paragraph(), after=4 if number < len(statements) else 9)
        paragraph.paragraph_format.left_indent = Cm(0.95)
        paragraph.paragraph_format.first_line_indent = Cm(-0.5)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.95))
        text_run(paragraph, f"{number}.\t{statement}", bold=True)
    cell_borders(declaration)

    signature = next((LOGOS / name for name in ("signature.png", "signature.jpg", "signature.jpeg")
                      if (LOGOS / name).exists()), None)
    signature_table = doc.add_table(rows=1, cols=2)
    set_column_widths(signature_table, [COVER_WIDTH_CM / 2] * 2)
    for index, label in enumerate(("Learner signature", "Tutor signature")):
        cell = signature_table.rows[0].cells[index]
        top = tight(cell.paragraphs[0], before=10)
        top.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if index == 0 and signature is not None:
            top.add_run().add_picture(str(signature), height=Cm(1.2))
        else:
            top.paragraph_format.space_before = Pt(36)
        for text, is_rule in (("_________________", True), (label, False), ("Date:", False)):
            paragraph = tight(cell.add_paragraph(), after=1)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text_run(paragraph, text, bold=not is_rule, italic=not is_rule)
        cell.paragraphs[-1].paragraph_format.space_after = Pt(8)
        cell_borders(cell)

    note_table = doc.add_table(rows=1, cols=1)
    set_column_widths(note_table, [COVER_WIDTH_CM])
    note = note_table.rows[0].cells[0]
    text_run(tight(note.paragraphs[0], before=9, after=9), "Note:")
    text_run(tight(note.add_paragraph(), after=22),
             "Assignments must be submitted in typed PDF format only; handwritten assignments are not accepted.")
    cell_borders(note)

    cover = doc.sections[0]
    cover.left_margin = cover.right_margin = Cm(COVER_MARGIN_CM)
    cover.top_margin = Cm(1.6)
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(body, side, Cm(MARGIN_CM))
    # Point the report section at the same "Page X of Y" footer explicitly,
    # rather than relying on inheritance from the cover section.
    footer_reference = doc.sections[0]._sectPr.find(qn("w:footerReference"))
    body._sectPr.insert(0, copy.deepcopy(footer_reference))


def contents_page(doc, entries: list[tuple[int, str, str]], pages: dict[str, int]) -> None:
    doc.add_paragraph("Table of Contents", style="Contents Title")
    for index, (level, text, bookmark) in enumerate(entries):
        paragraph = doc.add_paragraph(style=f"toc {level}")
        if index == 0:
            add_field_start(paragraph)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), bookmark)
        hyperlink.set(qn("w:history"), "1")
        for content in (text, "\t", str(pages.get(text, "")) or "?"):
            run = paragraph.add_run(content)
            set_run_font(run, BODY_FONT, 12)
            run.bold = level == 1
            hyperlink.append(run._element)
        paragraph._p.append(hyperlink)
        if index == len(entries) - 1:
            add_field_end(paragraph)


def add_field_start(paragraph) -> None:
    for kind, instruction in (("begin", None), ("instr", 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
        run = paragraph.add_run()
        if kind == "instr":
            text = OxmlElement("w:instrText")
            text.set(qn("xml:space"), "preserve")
            text.text = f" {instruction} "
            run._element.append(text)
        else:
            char = OxmlElement("w:fldChar")
            char.set(qn("w:fldCharType"), kind)
            run._element.append(char)


def add_field_end(paragraph) -> None:
    run = paragraph.add_run()
    char = OxmlElement("w:fldChar")
    char.set(qn("w:fldCharType"), "end")
    run._element.append(char)


# ============================================================ appendices

def commit_history() -> str:
    return subprocess.run(
        ["git", "log", "--reverse", "--date=format:%Y-%m-%d %H:%M", "--format=%h  %ad  %s"],
        cwd=WORK, capture_output=True, text=True, check=True).stdout.rstrip("\n")


def appendix_blocks() -> list[Block]:
    blocks = [Block("h1", "Appendix A: Git Repository"),
              Block("paragraph", f"The prototype, its tests, the CI workflow, the diagrams and the scripts that "
                                 f"build this report are in a public GitHub repository, shared with the "
                                 f"lecturer: **{REPOSITORY}**"),
              Block("paragraph", "To run the prototype, clone the repository and start it with Docker Compose. "
                                 "The README explains the endpoints and every environment variable."),
              Block("code", f"git clone {REPOSITORY}.git\ncd maltashop-microservices\n"
                            "docker compose up --build -d --wait\npython3 scripts/smoke_test.py"),
              Block("paragraph", "The commit history below, oldest first, shows the prototype being built step "
                                 "by step (it was taken when this report was built)."),
              Block("listing", commit_history()),
              Block("h1", "Appendix B: Prototype Source Code"),
              Block("paragraph", "The Compose file and the source files quoted in Section 6.1, in full. The "
                                 "Dockerfile and the CI workflow are in Listings 7 and 2.")]
    for title, path in APPENDIX_SOURCES:
        blocks.append(Block("h2", title))
        blocks.append(Block("listing", (WORK / path).read_text(encoding="utf-8").rstrip("\n"), target=path))
    return blocks


# ============================================================ main

def word_count() -> int:
    result = subprocess.run([sys.executable, str(ROOT / "wordcount.py")],
                            capture_output=True, text=True, check=True)
    return int(re.search(r"TOTAL_WORDS=(\d+)", result.stdout).group(1))


def main() -> int:
    body: list[Block] = []
    for path in sorted(CONTENT.glob("0*.md")):
        body += parse_markdown(path.read_text(encoding="utf-8"))
    references = parse_markdown((CONTENT / "90-references.md").read_text(encoding="utf-8"))
    appendices = appendix_blocks()

    entries = []
    bookmarks = {}
    for block in body + references + appendices:
        if block.kind in ("h1", "h2"):
            name = f"_Toc{len(entries) + 1:03d}"
            entries.append((int(block.kind[1]), block.text, name))
            bookmarks[block.text] = name

    pages = json.loads(TOC_PAGES.read_text()) if TOC_PAGES.exists() else {}
    count = word_count()

    doc = setup_document()
    cover_sheet(doc, count)
    contents_page(doc, entries, pages)
    counter = [0]
    write_blocks(doc, body, bookmarks, counter)

    for block in references:
        if block.kind == "h1":
            write_blocks(doc, [block], bookmarks, counter)
        elif block.kind == "paragraph":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(1.0)
            paragraph.paragraph_format.first_line_indent = Cm(-1.0)
            write_inline(paragraph, block.text)

    write_blocks(doc, appendices, bookmarks, counter)

    OUTPUT.mkdir(exist_ok=True)
    doc.save(DOCX_PATH)
    missing = [text for _, text, _ in entries if text not in pages]
    print(f"built: {DOCX_PATH.relative_to(WORK)}  ({count:,} assessed words, {len(entries)} contents entries"
          + (f", {len(missing)} without page numbers yet)" if missing else ")"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
